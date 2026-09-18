from datetime import date, datetime
from io import BytesIO
from flask import Blueprint, render_template, request, make_response
from flask_login import login_required, current_user
from extensions import db
from app.models import Barang, BarangMasuk, BarangMasukDetail, Permintaan, PermintaanDetail, Unit
from app.utils import admin_required, fmt_date_full

bp = Blueprint("laporan", __name__)


def _parse_date_range():
    today = date.today()
    start_str = request.args.get("start") or today.replace(day=1).isoformat()
    end_str = request.args.get("end") or today.isoformat()
    try:
        start = date.fromisoformat(start_str)
    except ValueError:
        start = today.replace(day=1)
    try:
        end = date.fromisoformat(end_str)
    except ValueError:
        end = today
    return start, end


# ===========================================================================
# Laporan Stok
# ===========================================================================
@bp.route("/stok")
@login_required
def stok():
    items = Barang.query.order_by(Barang.kode).all()
    total_nilai = sum(b.nilai_stok for b in items)
    return render_template("laporan/stok.html", items=items, total_nilai=total_nilai,
                           tanggal=date.today())


@bp.route("/stok/pdf")
@login_required
def stok_pdf():
    from xhtml2pdf import pisa
    items = Barang.query.order_by(Barang.kode).all()
    total_nilai = sum(b.nilai_stok for b in items)
    html = render_template("laporan/stok_pdf.html", items=items, total_nilai=total_nilai,
                           tanggal=date.today())
    result = BytesIO()
    pisa.CreatePDF(BytesIO(html.encode("utf-8")), result)
    resp = make_response(result.getvalue())
    resp.headers["Content-Type"] = "application/pdf"
    resp.headers["Content-Disposition"] = f"attachment; filename=laporan_stok_{date.today().isoformat()}.pdf"
    return resp


@bp.route("/stok/excel")
@login_required
def stok_excel():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Stok"
    ws.append(["No", "Kode", "Nama", "Kategori", "Satuan", "Harga Satuan", "Stok", "Stok Min", "Nilai Stok"])
    for i, b in enumerate(Barang.query.order_by(Barang.kode).all(), 1):
        ws.append([i, b.kode, b.nama, b.kategori.nama, b.satuan.kode,
                   float(b.harga_satuan), b.stok, b.stok_min, float(b.nilai_stok)])
    out = BytesIO()
    wb.save(out)
    resp = make_response(out.getvalue())
    resp.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    resp.headers["Content-Disposition"] = f"attachment; filename=laporan_stok_{date.today().isoformat()}.xlsx"
    return resp


# ===========================================================================
# Laporan Mutasi
# ===========================================================================
@bp.route("/mutasi")
@login_required
def mutasi():
    start, end = _parse_date_range()
    # Barang masuk
    masuk = BarangMasuk.query.filter(
        BarangMasuk.tanggal >= start, BarangMasuk.tanggal <= end
    ).order_by(BarangMasuk.tanggal).all()

    keluar = Permintaan.query.filter(
        Permintaan.tanggal >= start, Permintaan.tanggal <= end,
        Permintaan.status.in_([Permintaan.STATUS_DIPROSES, Permintaan.STATUS_SELESAI]),
    ).order_by(Permintaan.tanggal).all()

    return render_template("laporan/mutasi.html", masuk=masuk, keluar=keluar,
                           start=start, end=end)


@bp.route("/mutasi/pdf")
@login_required
def mutasi_pdf():
    from xhtml2pdf import pisa
    start, end = _parse_date_range()
    masuk = BarangMasuk.query.filter(
        BarangMasuk.tanggal >= start, BarangMasuk.tanggal <= end
    ).order_by(BarangMasuk.tanggal).all()
    keluar = Permintaan.query.filter(
        Permintaan.tanggal >= start, Permintaan.tanggal <= end,
        Permintaan.status.in_([Permintaan.STATUS_DIPROSES, Permintaan.STATUS_SELESAI]),
    ).order_by(Permintaan.tanggal).all()
    html = render_template("laporan/mutasi_pdf.html", masuk=masuk, keluar=keluar,
                           start=start, end=end)
    result = BytesIO()
    pisa.CreatePDF(BytesIO(html.encode("utf-8")), result)
    resp = make_response(result.getvalue())
    resp.headers["Content-Type"] = "application/pdf"
    resp.headers["Content-Disposition"] = f"attachment; filename=laporan_mutasi_{start}_{end}.pdf"
    return resp


@bp.route("/mutasi/excel")
@login_required
def mutasi_excel():
    from openpyxl import Workbook
    start, end = _parse_date_range()
    wb = Workbook()
    ws = wb.active
    ws.title = "Mutasi"
    ws.append(["Jenis", "Nomor", "Tanggal", "Supplier/Unit", "Barang", "Qty", "Satuan"])
    masuk = BarangMasuk.query.filter(
        BarangMasuk.tanggal >= start, BarangMasuk.tanggal <= end
    ).order_by(BarangMasuk.tanggal).all()
    for bm in masuk:
        for d in bm.details:
            ws.append(["Masuk", bm.no_faktur, bm.tanggal.isoformat(), bm.supplier.nama,
                       d.barang.nama, d.qty, d.barang.satuan.kode])
    keluar = Permintaan.query.filter(
        Permintaan.tanggal >= start, Permintaan.tanggal <= end,
        Permintaan.status.in_([Permintaan.STATUS_DIPROSES, Permintaan.STATUS_SELESAI]),
    ).order_by(Permintaan.tanggal).all()
    for p in keluar:
        for d in p.details:
            ws.append(["Keluar", p.no_permintaan, p.tanggal.isoformat(), p.unit.nama,
                       d.barang.nama, d.qty_diproses, d.barang.satuan.kode])
    out = BytesIO()
    wb.save(out)
    resp = make_response(out.getvalue())
    resp.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    resp.headers["Content-Disposition"] = f"attachment; filename=laporan_mutasi_{start}_{end}.xlsx"
    return resp


# ===========================================================================
# Rekap Pemakaian per Unit
# ===========================================================================
@bp.route("/pemakaian")
@login_required
def pemakaian():
    start, end = _parse_date_range()
    query = (
        db.session.query(
            Unit.kode.label("ukode"),
            Unit.nama.label("unama"),
            Barang.kode.label("bkode"),
            Barang.nama.label("bnama"),
            db.func.coalesce(db.func.sum(PermintaanDetail.qty_diproses), 0).label("total_qty"),
            db.func.coalesce(db.func.sum(PermintaanDetail.qty_diproses * Barang.harga_satuan), 0).label("total_nilai"),
        )
        .select_from(Permintaan)
        .join(PermintaanDetail, PermintaanDetail.header_id == Permintaan.id)
        .join(Barang, PermintaanDetail.barang_id == Barang.id)
        .join(Unit, Permintaan.unit_id == Unit.id)
        .filter(
            Permintaan.tanggal >= start,
            Permintaan.tanggal <= end,
            Permintaan.status.in_([Permintaan.STATUS_DIPROSES, Permintaan.STATUS_SELESAI]),
        )
        .group_by(Unit.kode, Unit.nama, Barang.kode, Barang.nama)
        .order_by(Unit.kode, db.desc("total_qty"))
    )
    rows = query.all()

    # Group by unit
    from collections import OrderedDict
    units_data = OrderedDict()
    for r in rows:
        key = r.ukode
        if key not in units_data:
            units_data[key] = {"kode": r.ukode, "nama": r.unama, "rows": [], "grand_qty": 0, "grand_nilai": 0}
        units_data[key]["rows"].append(r)
        units_data[key]["grand_qty"] += int(r.total_qty)
        units_data[key]["grand_nilai"] += float(r.total_nilai)

    return render_template("laporan/pemakaian.html", units_data=units_data,
                           start=start, end=end)


@bp.route("/pemakaian/pdf")
@login_required
def pemakaian_pdf():
    from xhtml2pdf import pisa
    start, end = _parse_date_range()
    query = (
        db.session.query(
            Unit.kode.label("ukode"),
            Unit.nama.label("unama"),
            Barang.kode.label("bkode"),
            Barang.nama.label("bnama"),
            db.func.coalesce(db.func.sum(PermintaanDetail.qty_diproses), 0).label("total_qty"),
            db.func.coalesce(db.func.sum(PermintaanDetail.qty_diproses * Barang.harga_satuan), 0).label("total_nilai"),
        )
        .select_from(Permintaan)
        .join(PermintaanDetail, PermintaanDetail.header_id == Permintaan.id)
        .join(Barang, PermintaanDetail.barang_id == Barang.id)
        .join(Unit, Permintaan.unit_id == Unit.id)
        .filter(
            Permintaan.tanggal >= start,
            Permintaan.tanggal <= end,
            Permintaan.status.in_([Permintaan.STATUS_DIPROSES, Permintaan.STATUS_SELESAI]),
        )
        .group_by(Unit.kode, Unit.nama, Barang.kode, Barang.nama)
        .order_by(Unit.kode, db.desc("total_qty"))
    )
    rows = query.all()
    from collections import OrderedDict
    units_data = OrderedDict()
    for r in rows:
        key = r.ukode
        if key not in units_data:
            units_data[key] = {"kode": r.ukode, "nama": r.unama, "rows": [], "grand_qty": 0, "grand_nilai": 0}
        units_data[key]["rows"].append(r)
        units_data[key]["grand_qty"] += int(r.total_qty)
        units_data[key]["grand_nilai"] += float(r.total_nilai)

    html = render_template("laporan/pemakaian_pdf.html", units_data=units_data,
                           start=start, end=end)
    result = BytesIO()
    pisa.CreatePDF(BytesIO(html.encode("utf-8")), result)
    resp = make_response(result.getvalue())
    resp.headers["Content-Type"] = "application/pdf"
    resp.headers["Content-Disposition"] = f"attachment; filename=rekap_pemakaian_{start}_{end}.pdf"
    return resp


@bp.route("/pemakaian/excel")
@login_required
def pemakaian_excel():
    from openpyxl import Workbook
    start, end = _parse_date_range()
    query = (
        db.session.query(
            Unit.kode, Unit.nama, Barang.kode, Barang.nama,
            db.func.coalesce(db.func.sum(PermintaanDetail.qty_diproses), 0).label("qty"),
            db.func.coalesce(db.func.sum(PermintaanDetail.qty_diproses * Barang.harga_satuan), 0).label("nilai"),
        )
        .select_from(Permintaan)
        .join(PermintaanDetail, PermintaanDetail.header_id == Permintaan.id)
        .join(Barang, PermintaanDetail.barang_id == Barang.id)
        .join(Unit, Permintaan.unit_id == Unit.id)
        .filter(
            Permintaan.tanggal >= start, Permintaan.tanggal <= end,
            Permintaan.status.in_([Permintaan.STATUS_DIPROSES, Permintaan.STATUS_SELESAI]),
        )
        .group_by(Unit.kode, Unit.nama, Barang.kode, Barang.nama)
        .order_by(Unit.kode, db.desc("qty"))
    )
    rows = query.all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap Pemakaian"
    ws.append(["Unit", "Kode Barang", "Nama Barang", "Total Qty", "Total Nilai (Rp)"])
    for r in rows:
        ws.append([r[1], r[2], r[3], int(r[4]), float(r[5])])
    out = BytesIO()
    wb.save(out)
    resp = make_response(out.getvalue())
    resp.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    resp.headers["Content-Disposition"] = f"attachment; filename=rekap_pemakaian_{start}_{end}.xlsx"
    return resp
