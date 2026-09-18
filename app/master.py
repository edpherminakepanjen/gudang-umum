from io import BytesIO
from openpyxl import Workbook, load_workbook
from flask import Blueprint, render_template, request, redirect, url_for, flash, make_response
from flask_login import login_required, current_user
from extensions import db
from app.models import Kategori, Satuan, Barang, Supplier, Unit, User, Permintaan
from app.utils import admin_required, log_audit

bp = Blueprint("master", __name__)


# ===========================================================================
# Kategori
# ===========================================================================
@bp.route("/kategori")
@login_required
@admin_required
def kategori_index():
    items = Kategori.query.order_by(Kategori.kode).all()
    satuans = Satuan.query.order_by(Satuan.kode).all()
    return render_template("master/kategori.html", items=items, satuans=satuans)


@bp.route("/kategori/create", methods=["POST"])
@login_required
@admin_required
def kategori_create():
    kode = (request.form.get("kode") or "").strip().upper()
    nama = (request.form.get("nama") or "").strip()
    keterangan = (request.form.get("keterangan") or "").strip()
    if not kode or not nama:
        flash("Kode dan nama kategori wajib diisi.", "error")
        return redirect(url_for("master.kategori_index"))
    if Kategori.query.filter_by(kode=kode).first():
        flash("Kode kategori sudah ada.", "error")
        return redirect(url_for("master.kategori_index"))
    k = Kategori(kode=kode, nama=nama, keterangan=keterangan)
    db.session.add(k)
    db.session.commit()
    log_audit("create", "kategori", target=f"{kode} - {nama}")
    flash("Kategori berhasil ditambahkan.", "success")
    return redirect(url_for("master.kategori_index"))


@bp.route("/kategori/<int:kid>/edit", methods=["POST"])
@login_required
@admin_required
def kategori_edit(kid):
    k = Kategori.query.get_or_404(kid)
    k.nama = (request.form.get("nama") or "").strip() or k.nama
    k.keterangan = (request.form.get("keterangan") or "").strip()
    db.session.commit()
    log_audit("update", "kategori", target=f"{k.kode} - {k.nama}")
    flash("Kategori diperbarui.", "success")
    return redirect(url_for("master.kategori_index"))


@bp.route("/kategori/<int:kid>/delete", methods=["POST"])
@login_required
@admin_required
def kategori_delete(kid):
    k = Kategori.query.get_or_404(kid)
    if k.barangs:
        flash("Tidak bisa menghapus kategori yang masih memiliki barang.", "error")
        return redirect(url_for("master.kategori_index"))
    db.session.delete(k)
    db.session.commit()
    log_audit("delete", "kategori", target=f"{k.kode} - {k.nama}")
    flash("Kategori dihapus.", "info")
    return redirect(url_for("master.kategori_index"))


# ===========================================================================
# Satuan
# ===========================================================================
@bp.route("/satuan")
@login_required
@admin_required
def satuan_index():
    return redirect(url_for("master.kategori_index"))


@bp.route("/satuan/create", methods=["POST"])
@login_required
@admin_required
def satuan_create():
    kode = (request.form.get("kode") or "").strip().lower()
    nama = (request.form.get("nama") or "").strip()
    if not kode or not nama:
        flash("Kode dan nama satuan wajib diisi.", "error")
        return redirect(url_for("master.satuan_index"))
    if Satuan.query.filter_by(kode=kode).first():
        flash("Kode satuan sudah ada.", "error")
        return redirect(url_for("master.satuan_index"))
    s = Satuan(kode=kode, nama=nama)
    db.session.add(s)
    db.session.commit()
    log_audit("create", "satuan", target=f"{kode} - {nama}")
    flash("Satuan ditambahkan.", "success")
    return redirect(url_for("master.satuan_index"))


@bp.route("/satuan/<int:sid>/edit", methods=["POST"])
@login_required
@admin_required
def satuan_edit(sid):
    s = Satuan.query.get_or_404(sid)
    s.nama = (request.form.get("nama") or "").strip() or s.nama
    db.session.commit()
    log_audit("update", "satuan", target=f"{s.kode} - {s.nama}")
    flash("Satuan diperbarui.", "success")
    return redirect(url_for("master.satuan_index"))


@bp.route("/satuan/<int:sid>/delete", methods=["POST"])
@login_required
@admin_required
def satuan_delete(sid):
    s = Satuan.query.get_or_404(sid)
    if s.barangs:
        flash("Tidak bisa menghapus satuan yang masih dipakai barang.", "error")
        return redirect(url_for("master.satuan_index"))
    db.session.delete(s)
    db.session.commit()
    log_audit("delete", "satuan", target=f"{s.kode} - {s.nama}")
    flash("Satuan dihapus.", "info")
    return redirect(url_for("master.satuan_index"))


# ===========================================================================
# Barang
# ===========================================================================
@bp.route("/barang")
@login_required
def barang_index():
    q = (request.args.get("q") or "").strip()
    kat = request.args.get("kategori", type=int)
    status_f = (request.args.get("status") or "aktif").strip()  # aktif | nonaktif | semua

    query = Barang.query
    if q:
        query = query.filter(
            db.or_(Barang.nama.ilike(f"%{q}%"), Barang.kode.ilike(f"%{q}%"))
        )
    if kat:
        query = query.filter_by(kategori_id=kat)
    if status_f == "aktif":
        query = query.filter_by(is_aktif=True)
    elif status_f == "nonaktif":
        query = query.filter_by(is_aktif=False)
    # status_f == "semua" -> tidak difilter

    items = query.order_by(Barang.kode).all()
    kategoris = Kategori.query.order_by(Kategori.nama).all()
    satuans = Satuan.query.order_by(Satuan.kode).all()
    return render_template("master/barang.html", items=items, kategoris=kategoris,
                           satuans=satuans, q=q, kat=kat, status_f=status_f)


@bp.route("/barang/create", methods=["POST"])
@login_required
@admin_required
def barang_create():
    kode = (request.form.get("kode") or "").strip().upper()
    nama = (request.form.get("nama") or "").strip()
    kategori_id = request.form.get("kategori_id", type=int)
    satuan_id = request.form.get("satuan_id", type=int)
    harga = request.form.get("harga_satuan", type=float) or 0
    stok = request.form.get("stok", type=int) or 0
    stok_min = request.form.get("stok_min", type=int) or 0
    keterangan = (request.form.get("keterangan") or "").strip()

    if not kode or not nama or not kategori_id or not satuan_id:
        flash("Kode, nama, kategori, dan satuan wajib diisi.", "error")
        return redirect(url_for("master.barang_index"))
    if Barang.query.filter_by(kode=kode).first():
        flash("Kode barang sudah ada.", "error")
        return redirect(url_for("master.barang_index"))

    b = Barang(kode=kode, nama=nama, kategori_id=kategori_id, satuan_id=satuan_id,
               harga_satuan=harga, stok=stok, stok_min=stok_min, keterangan=keterangan)
    db.session.add(b)
    db.session.commit()
    log_audit("create", "barang", target=f"{kode} - {nama}")
    flash("Barang ditambahkan.", "success")
    return redirect(url_for("master.barang_index"))


@bp.route("/barang/<int:bid>/edit", methods=["POST"])
@login_required
@admin_required
def barang_edit(bid):
    b = Barang.query.get_or_404(bid)
    b.nama = (request.form.get("nama") or "").strip() or b.nama
    b.kategori_id = request.form.get("kategori_id", type=int) or b.kategori_id
    b.satuan_id = request.form.get("satuan_id", type=int) or b.satuan_id
    b.harga_satuan = request.form.get("harga_satuan", type=float) or 0
    b.stok = request.form.get("stok", type=int) or 0
    b.stok_min = request.form.get("stok_min", type=int) or 0
    b.keterangan = (request.form.get("keterangan") or "").strip()
    db.session.commit()
    log_audit("update", "barang", target=f"{b.kode} - {b.nama}")
    flash("Barang diperbarui.", "success")
    return redirect(url_for("master.barang_index"))


@bp.route("/barang/<int:bid>/toggle", methods=["POST"])
@login_required
@admin_required
def barang_toggle(bid):
    b = Barang.query.get_or_404(bid)
    b.is_aktif = not b.is_aktif
    db.session.commit()
    aksi = "activate" if b.is_aktif else "deactivate"
    log_audit(aksi, "barang", target=f"{b.kode} - {b.nama}")
    if b.is_aktif:
        flash(f"Barang '{b.nama}' diaktifkan kembali.", "success")
    else:
        flash(f"Barang '{b.nama}' dinonaktifkan. Barang tidak akan muncul di transaksi baru, tapi riwayatnya tetap aman.", "info")
    return redirect(url_for("master.barang_index", status=request.args.get("status", "aktif")))

@bp.route("/barang/import/template")
@login_required
@admin_required
def barang_import_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Barang"
    ws.append(["Kode", "Nama", "Kategori", "Satuan", "Harga Satuan", "Stok Awal", "Stok Minimum", "Keterangan"])
    ws.append(["ATK-100", "Contoh: Spidol Whiteboard", "ATK", "pcs", 8000, 20, 5, "Opsional"])

    ws_ref = wb.create_sheet("Referensi Kategori & Satuan")
    ws_ref.append(["Kode Kategori", "Nama Kategori"])
    for k in Kategori.query.order_by(Kategori.kode).all():
        ws_ref.append([k.kode, k.nama])
    ws_ref.append([])
    ws_ref.append(["Kode Satuan", "Nama Satuan"])
    for s in Satuan.query.order_by(Satuan.kode).all():
        ws_ref.append([s.kode, s.nama])

    out = BytesIO()
    wb.save(out)
    resp = make_response(out.getvalue())
    resp.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    resp.headers["Content-Disposition"] = "attachment; filename=template_import_barang.xlsx"
    return resp


@bp.route("/barang/import", methods=["GET", "POST"])
@login_required
@admin_required
def barang_import():
    if request.method == "POST":
        file = request.files.get("file")
        if not file or not file.filename.lower().endswith((".xlsx", ".xls")):
            flash("Upload file Excel (.xlsx) yang valid.", "error")
            return redirect(url_for("master.barang_import"))

        try:
            wb = load_workbook(file, data_only=True)
            ws = wb["Barang"] if "Barang" in wb.sheetnames else wb.active
        except Exception:
            flash("File Excel tidak bisa dibaca. Pastikan formatnya benar.", "error")
            return redirect(url_for("master.barang_import"))

        kategoris = {k.kode.upper(): k for k in Kategori.query.all()}
        satuans = {s.kode.lower(): s for s in Satuan.query.all()}

        berhasil, dilewati, error_rows = 0, 0, []
        rows = list(ws.iter_rows(min_row=2, values_only=True))

        for i, row in enumerate(rows, start=2):
            if not row or not any(row):
                continue
            kode, nama, kat_kode, sat_kode, harga, stok, stok_min, ket = (list(row) + [None] * 8)[:8]

            kode = str(kode).strip().upper() if kode else ""
            nama = str(nama).strip() if nama else ""
            kat_kode = str(kat_kode).strip().upper() if kat_kode else ""
            sat_kode = str(sat_kode).strip().lower() if sat_kode else ""

            if not kode or not nama or not kat_kode or not sat_kode:
                error_rows.append(f"Baris {i}: kode/nama/kategori/satuan tidak lengkap")
                continue
            if Barang.query.filter_by(kode=kode).first():
                dilewati += 1
                continue
            if kat_kode not in kategoris:
                error_rows.append(f"Baris {i}: kategori '{kat_kode}' tidak ditemukan di sistem")
                continue
            if sat_kode not in satuans:
                error_rows.append(f"Baris {i}: satuan '{sat_kode}' tidak ditemukan di sistem")
                continue

            try:
                harga_val = float(harga) if harga not in (None, "") else 0
                stok_val = int(stok) if stok not in (None, "") else 0
                stok_min_val = int(stok_min) if stok_min not in (None, "") else 0
            except (TypeError, ValueError):
                error_rows.append(f"Baris {i}: harga/stok/stok minimum harus berupa angka")
                continue

            b = Barang(
                kode=kode, nama=nama,
                kategori_id=kategoris[kat_kode].id,
                satuan_id=satuans[sat_kode].id,
                harga_satuan=harga_val, stok=stok_val, stok_min=stok_min_val,
                keterangan=(str(ket).strip() if ket else ""),
            )
            db.session.add(b)
            berhasil += 1

        db.session.commit()
        log_audit("import", "barang", target=f"{berhasil} barang",
                   detail=f"dilewati: {dilewati}, error: {len(error_rows)}")

        flash(
            f"Import selesai: {berhasil} barang berhasil ditambahkan, {dilewati} dilewati (kode sudah ada), {len(error_rows)} baris error.",
            "success" if berhasil else "warning",
        )
        return render_template("master/barang_import.html", error_rows=error_rows, done=True,
                               berhasil=berhasil, dilewati=dilewati)

    return render_template("master/barang_import.html", done=False)

# ===========================================================================
# Supplier
# ===========================================================================
@bp.route("/supplier")
@login_required
@admin_required
def supplier_index():
    items = Supplier.query.order_by(Supplier.kode).all()
    return render_template("master/supplier.html", items=items)


@bp.route("/supplier/create", methods=["POST"])
@login_required
@admin_required
def supplier_create():
    kode = (request.form.get("kode") or "").strip().upper()
    nama = (request.form.get("nama") or "").strip()
    if not kode or not nama:
        flash("Kode dan nama supplier wajib diisi.", "error")
        return redirect(url_for("master.supplier_index"))
    if Supplier.query.filter_by(kode=kode).first():
        flash("Kode supplier sudah ada.", "error")
        return redirect(url_for("master.supplier_index"))
    s = Supplier(
        kode=kode, nama=nama,
        kontak=(request.form.get("kontak") or "").strip(),
        alamat=(request.form.get("alamat") or "").strip(),
        telepon=(request.form.get("telepon") or "").strip(),
        email=(request.form.get("email") or "").strip(),
        keterangan=(request.form.get("keterangan") or "").strip(),
    )
    db.session.add(s)
    db.session.commit()
    log_audit("create", "supplier", target=f"{kode} - {nama}")
    flash("Supplier ditambahkan.", "success")
    return redirect(url_for("master.supplier_index"))


@bp.route("/supplier/<int:sid>/edit", methods=["POST"])
@login_required
@admin_required
def supplier_edit(sid):
    s = Supplier.query.get_or_404(sid)
    s.nama = (request.form.get("nama") or "").strip() or s.nama
    s.kontak = (request.form.get("kontak") or "").strip()
    s.alamat = (request.form.get("alamat") or "").strip()
    s.telepon = (request.form.get("telepon") or "").strip()
    s.email = (request.form.get("email") or "").strip()
    s.keterangan = (request.form.get("keterangan") or "").strip()
    db.session.commit()
    log_audit("update", "supplier", target=f"{s.kode} - {s.nama}")
    flash("Supplier diperbarui.", "success")
    return redirect(url_for("master.supplier_index"))


@bp.route("/supplier/<int:sid>/delete", methods=["POST"])
@login_required
@admin_required
def supplier_delete(sid):
    s = Supplier.query.get_or_404(sid)
    if s.barang_masuk_list:
        flash("Tidak bisa menghapus supplier dengan riwayat barang masuk.", "error")
        return redirect(url_for("master.supplier_index"))
    db.session.delete(s)
    db.session.commit()
    log_audit("delete", "supplier", target=f"{s.kode} - {s.nama}")
    flash("Supplier dihapus.", "info")
    return redirect(url_for("master.supplier_index"))


# ===========================================================================
# Unit
# ===========================================================================
@bp.route("/unit")
@login_required
def unit_index():
    if not current_user.is_admin:
        return redirect(url_for("dashboard.index"))
    items = Unit.query.order_by(Unit.kode).all()
    return render_template("master/unit.html", items=items)


@bp.route("/unit/create", methods=["POST"])
@login_required
@admin_required
def unit_create():
    kode = (request.form.get("kode") or "").strip().upper()
    nama = (request.form.get("nama") or "").strip()
    if not kode or not nama:
        flash("Kode dan nama unit wajib diisi.", "error")
        return redirect(url_for("master.unit_index"))
    if Unit.query.filter_by(kode=kode).first():
        flash("Kode unit sudah ada.", "error")
        return redirect(url_for("master.unit_index"))
    u = Unit(kode=kode, nama=nama,
             keterangan=(request.form.get("keterangan") or "").strip())
    db.session.add(u)
    db.session.commit()
    log_audit("create", "unit", target=f"{kode} - {nama}")
    flash("Unit ditambahkan.", "success")
    return redirect(url_for("master.unit_index"))


@bp.route("/unit/<int:uid>/edit", methods=["POST"])
@login_required
@admin_required
def unit_edit(uid):
    u = Unit.query.get_or_404(uid)
    u.nama = (request.form.get("nama") or "").strip() or u.nama
    u.keterangan = (request.form.get("keterangan") or "").strip()
    db.session.commit()
    log_audit("update", "unit", target=f"{u.kode} - {u.nama}")
    flash("Unit diperbarui.", "success")
    return redirect(url_for("master.unit_index"))


@bp.route("/unit/<int:uid>/delete", methods=["POST"])
@login_required
@admin_required
def unit_delete(uid):
    u = Unit.query.get_or_404(uid)
    ada_permintaan = Permintaan.query.filter_by(unit_id=u.id).first() is not None
    if u.users or ada_permintaan:
        flash("Tidak bisa menghapus unit yang masih memiliki user atau permintaan.", "error")
        return redirect(url_for("master.unit_index"))
    db.session.delete(u)
    db.session.commit()
    log_audit("delete", "unit", target=f"{u.kode} - {u.nama}")
    flash("Unit dihapus.", "info")
    return redirect(url_for("master.unit_index"))
